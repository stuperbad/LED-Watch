"""Builds Watch-Power-Calculator.xlsx (live formulas). Re-run after editing defaults here."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference, Series

ARIAL = "Arial"
F_NORM = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_SEC = PatternFill("solid", fgColor="D9E2F3")
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def put(ws, ref, v, font=F_NORM, fmt=None, fill=None, bold=False, wrap=False, align=None):
    c = ws[ref]
    c.value = v
    c.font = Font(name=font.name, size=font.size, bold=bold or font.bold, italic=font.italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return c


def inp(ws, ref, v, fmt=None):
    c = put(ws, ref, v, F_INPUT, fmt, FILL_IN)
    c.border = BOX
    return c


def head(ws, row, labels, col0=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=t)
        c.font = F_HEAD; c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX


def section(ws, row, text, ncols=4):
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = FILL_SEC
    c = ws.cell(row=row, column=1, value=text); c.font = F_BOLD; c.fill = FILL_SEC


# =============================================================== Inputs
ws = wb.active; ws.title = "Inputs"
ws.column_dimensions["A"].width = 54; ws.column_dimensions["B"].width = 20; ws.column_dimensions["C"].width = 10; ws.column_dimensions["D"].width = 95
put(ws, "A1", "Watch Power Calculator - change the yellow cells (blue text)", F_TITLE)
put(ws, "A2", "Yellow cells with blue text are inputs you can change. Every other number in this workbook is a formula and updates automatically. "
              "VERIFIED = from a datasheet; ESTIMATE = my assumption, so please change it if you know better.", F_NOTE, wrap=True)
ws.merge_cells("A2:D2"); ws.row_dimensions[2].height = 28

section(ws, 4, "1. Battery")
rows = [
    (5, "Battery nominal voltage", 3.7, "V", "0.00", "LiPo nominal voltage (3.0-4.2 V in use). ESTIMATE of the average."),
    (6, "Battery capacity", 100, "mAh", "0", "Your choice. Candidate protected LiPos: 35, 50, 80 mAh (Renata) up to ~150 mAh."),
]
for r, lab, v, u, fmt, note in rows:
    put(ws, f"A{r}", lab); inp(ws, f"B{r}", v, fmt); put(ws, f"C{r}", u); put(ws, f"D{r}", note, F_NOTE)
put(ws, "A7", "Battery energy"); put(ws, "B7", "=B5*B6", fmt="0.0", fill=FILL_OUT); put(ws, "C7", "mWh"); put(ws, "D7", "Formula: voltage x capacity.", F_NOTE)
put(ws, "A8", "Battery self-discharge per month"); inp(ws, "B8", 0.02, "0.0%"); put(ws, "C8", "fraction"); put(ws, "D8", "ESTIMATE: 2%/month (LiPo is typically 1-5%).", F_NOTE)

section(ws, 10, "2. White LED electrical (one series resistor value for all columns)")
put(ws, "A11", "Series resistor in each column"); inp(ws, "B11", 150, "0"); put(ws, "C11", "ohm"); put(ws, "D11", "PLACEHOLDER in the schematic (R1-R12 = 150 ohm). White LEDs need ~2.7+ V, so this value sets the brightness.", F_NOTE)
put(ws, "A12", "Voltage lost in the MCU pins (high-side + low-side)"); inp(ws, "B12", 0.3, "0.00"); put(ws, "C12", "V"); put(ws, "D12", "ESTIMATE; check the AVR128DB data sheet's output-voltage curves.", F_NOTE)
put(ws, "A13", "OPTIONAL: force an LED on-current (leave blank to calculate)"); inp(ws, "B13", None, "0.00"); put(ws, "C13", "mA"); put(ws, "D13", "If you type a number here it replaces the calculated on-current for every LED.", F_NOTE)
put(ws, "D14", "On-current = (battery V - LED Vf - pin drop) / resistor. Vf per LED type is on the 'LED types' tab.", F_NOTE)

section(ws, 15, "3. Multiplexing and PWM dimming")
put(ws, "A16", "Rows scanned in turn while the display is on"); inp(ws, "B16", 3, "0"); put(ws, "C16", "rows")
put(ws, "D16", "3 when you show an hour + a minute + a second LED (they are always in different rows). Each LED then gets 1/3 of the time.", F_NOTE)
for r, lab in ((17, "PWM duty - Hour LED"), (18, "PWM duty - Minute LED"), (19, "PWM duty - Second LED")):
    put(ws, f"A{r}", lab); inp(ws, f"B{r}", 0.2, "0%"); put(ws, f"C{r}", "of slot")
put(ws, "D17", "PWM duty = how much of its time slot the LED is actually on. 100% = full brightness for the slot, 20% = dim.", F_NOTE)
put(ws, "D18", "Brightness and energy both follow the AVERAGE current, so halving the duty halves both.", F_NOTE)
put(ws, "D19", "Default 20% gives ~0.3 mA average per LED (dim but visible indoors).", F_NOTE)

section(ws, 21, "4. The time being shown, and how often")
put(ws, "A22", "Hour shown (1-12)"); inp(ws, "B22", 10, "0")
put(ws, "A23", "Minute shown (0-59)"); inp(ws, "B23", 8, "0")
put(ws, "A24", "Second shown (0-59)"); inp(ws, "B24", 42, "0")
put(ws, "D22", "These three pick which LEDs are lit on the 'Every LED' tab.", F_NOTE)
put(ws, "A25", "Seconds the display stays on after each wake"); inp(ws, "B25", 3, "0.0"); put(ws, "C25", "s")
put(ws, "A26", "Wakes per day (gesture or button)"); inp(ws, "B26", 60, "0"); put(ws, "C26", "per day"); put(ws, "D26", "20 = rarely, 60 = a few times an hour, 200 = checking constantly.", F_NOTE)
put(ws, "A27", "MCU current while the display is up"); inp(ws, "B27", 1.5, "0.00"); put(ws, "C27", "mA"); put(ws, "D27", "ESTIMATE; depends on clock speed (AVR128DB data sheet gives the curve).", F_NOTE)

section(ws, 29, "5. Always-on draw (the watch waiting to be woken)")
put(ws, "A30", "MCU deep sleep with 32.768 kHz clock running"); inp(ws, "B30", 3.3, "0.0"); put(ws, "C30", "uA"); put(ws, "D30", "Derived from the AVR128DB data sheet (power-down 0.65 uA + RTC/crystal ~2.6 uA); not a printed figure.", F_NOTE)
put(ws, "A31", "3.3 V regulator quiescent current"); inp(ws, "B31", 1.6, "0.0"); put(ws, "C31", "uA"); put(ws, "D31", "VERIFIED: MCP1700 data sheet (Variant A). The BQ25570 buck (Variant B) is similar or lower.", F_NOTE)
put(ws, "A32", "How the watch is woken (pick from the list)"); inp(ws, "B32", "LIS2DW12 accelerometer wake")
put(ws, "D32", "Button only = nothing running. BMI270 = the 6-axis sensor in wake-on-motion. BMA400 = a simpler accelerometer.", F_NOTE)
put(ws, "A33", "Custom wake sensor current (used only if you pick Custom)"); inp(ws, "B33", 5, "0.0"); put(ws, "C33", "uA")
put(ws, "A34", "Custom wake sensor supply voltage"); inp(ws, "B34", 3.3, "0.0"); put(ws, "C34", "V")

section(ws, 36, "6. Solar (optional)")
put(ws, "A37", "Solar cell (pick from the list)"); inp(ws, "B37", "KXOB25-07X1TF")
put(ws, "A38", "Wearer profile (pick from the list)"); inp(ws, "B38", "Commuter")
put(ws, "A39", "Wrist factor (face tilted, sleeve shade)"); inp(ws, "B39", 0.5, "0.00"); put(ws, "D39", "ESTIMATE (research report: wrist reads under half of eye-level light).", F_NOTE)
put(ws, "A40", "Optical / area factor (bezel, glass, LED rings)"); inp(ws, "B40", 0.85, "0.00"); put(ws, "D40", "ESTIMATE.", F_NOTE)
put(ws, "A41", "Harvester IC efficiency"); inp(ws, "B41", 0.75, "0.00"); put(ws, "D41", "ESTIMATE; real harvesters do worse at microwatt levels.", F_NOTE)
put(ws, "A42", "Charge efficiency into the battery"); inp(ws, "B42", 0.95, "0.00"); put(ws, "D42", "ESTIMATE.", F_NOTE)

section(ws, 44, "7. Special mode: fluid / gravity / tilt game (optional, a minute or two a day)")
put(ws, "A45", "Minutes of special mode per day"); inp(ws, "B45", 2, "0.0"); put(ws, "C45", "min")
put(ws, "A46", "LEDs lit at once (average) while playing"); inp(ws, "B46", 40, "0"); put(ws, "C46", "LEDs"); put(ws, "D46", "A 'liquid' filling about a third of the face = ~40 LEDs.", F_NOTE)
put(ws, "A47", "Rows scanned in special mode"); inp(ws, "B47", 11, "0"); put(ws, "C47", "rows"); put(ws, "D47", "11 = every row takes a turn, so each LED is on only 1/11 of the time.", F_NOTE)
put(ws, "A48", "PWM duty in special mode"); inp(ws, "B48", 0.7, "0%"); put(ws, "D48", "70% gives about the same brightness per LED as the clock mode (0.3 mA average); lower it to save energy.", F_NOTE)
put(ws, "A49", "Most LEDs lit in ONE row at the same instant (worst case)"); inp(ws, "B49", 12, "0"); put(ws, "C49", "LEDs"); put(ws, "D49", "A full row has 12 LEDs. All of a row's LEDs draw current through that one MCU pin.", F_NOTE)
put(ws, "A50", "Most current I will allow through one MCU pin"); inp(ws, "B50", 20, "0"); put(ws, "C50", "mA"); put(ws, "D50", "Conservative ESTIMATE. A web search suggests an absolute maximum around 40 mA per pin, but I could not read the AVR128DB data sheet's own table (UNVERIFIED).", F_NOTE)
put(ws, "A51", "Motion sensor mode during the game (pick from the list)"); inp(ws, "B51", "LIS2DW12 normal mode")
put(ws, "A52", "Custom sensor current during the game (used only if you pick Custom)"); inp(ws, "B52", 100, "0.0"); put(ws, "C52", "uA")
put(ws, "A53", "MCU current while running the game"); inp(ws, "B53", 3, "0.00"); put(ws, "C53", "mA"); put(ws, "D53", "ESTIMATE: physics + drawing needs a faster clock than the clock mode.", F_NOTE)
for r in range(1, 54):
    for col in "ABCD":
        c = ws[f"{col}{r}"]
        if c.font is None or c.font.name != ARIAL:
            c.font = F_NORM

# =============================================================== Lists
wl = wb.create_sheet("Lists")
head(wl, 1, ["Wake method", "Current (uA)", "Voltage (V)", "Source"])
lst = [("Button only (no motion sensor)", 0, 3.3, "Nothing running while waiting."),
       ("BMI270 gesture wake", 13, 3.3, "VERIFIED by the research agent (BMI270 wake-on-motion ~13 uA)."),
       ("BMA400 accelerometer wake", 0.85, 3.3, "VERIFIED by the research agent (BMA400 low-power mode 850 nA)."),
       ("LIS2DW12 accelerometer wake", 1.0, 3.3, "ESTIMATE (UNVERIFIED): ST datasheet not read."),
       ("Custom", "=Inputs!B33", "=Inputs!B34", "Uses the custom cells on the Inputs tab.")]
for i, (n, a, v, s) in enumerate(lst, start=2):
    put(wl, f"A{i}", n); put(wl, f"B{i}", a, F_LINK if isinstance(a, str) else F_NORM); put(wl, f"C{i}", v, F_LINK if isinstance(v, str) else F_NORM); put(wl, f"D{i}", s, F_NOTE)
wl.column_dimensions["A"].width = 34; wl.column_dimensions["D"].width = 70
dv1 = DataValidation(type="list", formula1="=Lists!$A$2:$A$6", allow_blank=False); ws.add_data_validation(dv1); dv1.add("B32")


head(wl, 1, ["Game-mode sensor setting", "Current (uA)", "Source"], col0=6)
sp = [("BMA400 normal mode", 14.5, "VERIFIED by the research agent (BMA400 normal mode 14.5 uA)."),
      ("BMI270 accelerometer only, normal", 210, "VERIFIED: BMI270 datasheet Table 6 (VDD 1.8 V test): accel-only normal mode 210 uA."),
      ("BMI270 accelerometer + gyroscope, normal", 685, "VERIFIED: BMI270 datasheet Table 6: A+G normal mode 685 uA."),
      ("LIS2DW12 normal mode", 50, "ESTIMATE (UNVERIFIED): ST datasheet not read."),
      ("Custom", "=Inputs!B52", "Uses the custom cell on the Inputs tab.")]
for i, (n, a, src) in enumerate(sp, start=2):
    put(wl, f"F{i}", n); put(wl, f"G{i}", a, F_LINK if isinstance(a, str) else F_NORM); put(wl, f"H{i}", src, F_NOTE)
wl.column_dimensions["F"].width = 40; wl.column_dimensions["H"].width = 70
dv4 = DataValidation(type="list", formula1="=Lists!$F$2:$F$6", allow_blank=False); ws.add_data_validation(dv4); dv4.add("B51")

# =============================================================== LED types
wt = wb.create_sheet("LED types")
put(wt, "A1", "LED types: electrical model and estimated brightness", F_TITLE)
put(wt, "A2", "One row per LED type. The 'Every LED' tab looks these up. Brightness is estimated as proportional to AVERAGE current (ESTIMATE); roughly 10-20 mcd looks clearly lit indoors (rule of thumb, ESTIMATE).", F_NOTE, wrap=True)
wt.merge_cells("A2:N2"); wt.row_dimensions[2].height = 30
head(wt, 4, ["Type", "Count in watch", "Package", "LED forward voltage Vf (V)", "PWM duty", "On-current when its slot is on (mA)", "Time-share (1 / rows)",
             "Average current while lit (mA)", "Average battery power while lit (mW)", "Instantaneous power when on (mW)",
             "Reference brightness (mcd)", "At reference current (mA)", "Estimated brightness while lit (mcd)", "Share of power that becomes LED (Vf / battery V)"])
wt.row_dimensions[4].height = 58
types = [("Hour", 12, "0603 (KT-0603W)", 250, 5), ("Minute", 60, "0402 (XL-1005UWC)", 600, 20), ("Second", 60, "0402 (XL-1005UWC)", 600, 20)]
for i, (t, n, pk, mcd, iref) in enumerate(types):
    r = 5 + i
    put(wt, f"A{r}", t, bold=True); put(wt, f"B{r}", n, fmt="0"); put(wt, f"C{r}", pk)
    inp(wt, f"D{r}", 2.7, "0.00")
    put(wt, f"E{r}", f"=Inputs!B{17 + i}", F_LINK, "0%")
    put(wt, f"F{r}", f"=IF(ISNUMBER(Inputs!$B$13),Inputs!$B$13,MAX(0,(Inputs!$B$5-D{r}-Inputs!$B$12)/Inputs!$B$11*1000))", fmt="0.000")
    put(wt, f"G{r}", "=1/Inputs!$B$16", fmt="0.000")
    put(wt, f"H{r}", f"=F{r}*G{r}*E{r}", fmt="0.000", fill=FILL_OUT)
    put(wt, f"I{r}", f"=Inputs!$B$5*H{r}", fmt="0.000", fill=FILL_OUT)
    put(wt, f"J{r}", f"=Inputs!$B$5*F{r}", fmt="0.00")
    inp(wt, f"K{r}", mcd, "0"); inp(wt, f"L{r}", iref, "0.0")
    put(wt, f"M{r}", f"=K{r}*H{r}/L{r}", fmt="0.0", fill=FILL_OUT)
    put(wt, f"N{r}", f"=IF(Inputs!$B$5>0,D{r}/Inputs!$B$5,0)", fmt="0%")
put(wt, "A9", "Sources: Vf 2.6-3.3 V and brightness 350-1000 mcd @ 20 mA (0402 XL-1005UWC) and 145-360 mcd @ 5 mA (0603/0805 KT-0603W/-0805W) are from the LCSC data sheets via the JLCPCB research report (jlc-alternatives.md). The 2.7 V Vf at low current and the mid-range brightness values are ESTIMATES.", F_NOTE, wrap=True)
wt.merge_cells("A9:N9"); wt.row_dimensions[9].height = 30
for col, w in zip("ABCDEFGHIJKLMN", (10, 10, 20, 14, 10, 16, 12, 16, 16, 16, 14, 14, 16, 18)):
    wt.column_dimensions[col].width = w

# =============================================================== Every LED
we = wb.create_sheet("Every LED")
put(we, "A1", "Every one of the 132 LEDs: position, matrix wiring, and power draw", F_TITLE)
put(we, "A2", "LEDs lit right now:", F_BOLD); put(we, "C2", "=SUM(K5:K136)", fmt="0", fill=FILL_OUT)
put(we, "D2", "Their total battery power right now (mW):", F_BOLD); put(we, "H2", "=SUM(M5:M136)", fmt="0.000", fill=FILL_OUT)
put(we, "A3", "Change the shown time on the Inputs tab (rows 22-24) to light a different set of 3 LEDs.", F_NOTE)
head(we, 4, ["LED", "Type", "Position (hour 1-12, or minute/second 0-59)", "Clock angle (deg, clockwise from 12)", "Matrix row", "Matrix column",
             "On-current (mA)", "Average current when lit (mA)", "Average battery power when lit (mW)", "Instantaneous power in its slot (mW)",
             "Lit now? (1 = yes)", "Current now (mA)", "Power now (mW)", "Energy for one display event (mWh)"])
we.row_dimensions[4].height = 62
rows_led = [(f"H{n}", "Hour", n) for n in range(1, 13)] + [(f"M{n}", "Minute", n) for n in range(60)] + [(f"S{n}", "Second", n) for n in range(60)]
for i, (ref, tp, pos) in enumerate(rows_led):
    r = 5 + i
    put(we, f"A{r}", ref); put(we, f"B{r}", tp); put(we, f"C{r}", pos, fmt="0")
    put(we, f"D{r}", f'=IF(B{r}="Hour",MOD(C{r},12)*30,C{r}*6)', fmt="0")
    put(we, f"E{r}", f'=IF(B{r}="Hour",0,IF(B{r}="Minute",1+MOD(C{r},5),6+MOD(C{r},5)))', fmt="0")
    put(we, f"F{r}", f'=IF(B{r}="Hour",MOD(C{r},12),INT(C{r}/5))', fmt="0")
    for col, src in (("G", "F"), ("H", "H"), ("I", "I"), ("J", "J")):
        put(we, f"{col}{r}", f"=INDEX('LED types'!${src}$5:${src}$7,MATCH($B{r},'LED types'!$A$5:$A$7,0))", fmt="0.000")
    put(we, f"K{r}", f'=IF(OR(AND(B{r}="Hour",C{r}=Inputs!$B$22),AND(B{r}="Minute",C{r}=Inputs!$B$23),AND(B{r}="Second",C{r}=Inputs!$B$24)),1,0)', fmt="0")
    put(we, f"L{r}", f"=K{r}*H{r}", fmt="0.000"); put(we, f"M{r}", f"=K{r}*I{r}", fmt="0.000")
    put(we, f"N{r}", f"=K{r}*I{r}*Inputs!$B$25/3600", fmt="0.00000")
for col, w in zip("ABCDEFGHIJKLMN", (8, 9, 22, 18, 10, 10, 12, 16, 18, 18, 12, 12, 12, 18)):
    we.column_dimensions[col].width = w
we.freeze_panes = "A5"

# =============================================================== Solar
sw = wb.create_sheet("Solar")
put(sw, "A1", "Solar: light units, cell power, and daily harvest", F_TITLE)
put(sw, "A2", "Unit chain: lux -> W/m2 (divide by lux-per-W/m2) -> mW/cm2 (divide by 10) -> suns (divide W/m2 by 1000). Indoor LED/fluorescent ~300 lux per W/m2, daylight ~120 (ESTIMATES, see solar-watch-evidence.md).", F_NOTE, wrap=True)
sw.merge_cells("A2:M2"); sw.row_dimensions[2].height = 28
head(sw, 4, ["Light condition", "Lux", "Lux per W/m2", "W/m2", "mW/cm2", "Suns", "Mono-Si low-light derating (vs 1 sun)", "Office hours/day", "Commuter hours/day", "Outdoor hours/day",
             "AM-1606C power (mW)", "KXOB25-07X1TF power (mW)", "KXOB081K06TF power (mW)"])
sw.row_dimensions[4].height = 58
bands = [("Dim indoor (home)", 200, 300, "=4.4/18.4", 6, 5, 3), ("Desk", 500, 300, "=7/18.4", 6, 4, 1), ("Window", 1000, 120, "=11/18.4", 1, 1, 0),
         ("Overcast / shade", 5000, 120, "=15/18.4", 0.7, 1.5, 3), ("Bright shade", 20000, 120, "=17/18.4", 0, 0, 3),
         ("Partial sun", 30000, 120, "=17.5/18.4", 0.3, 1.0, 0), ("Sunny", 50000, 120, "=18/18.4", 0, 0, 2)]
for i, (n, lux, lpw, der, ho, hc, hout) in enumerate(bands):
    r = 5 + i
    put(sw, f"A{r}", n); inp(sw, f"B{r}", lux, "#,##0"); inp(sw, f"C{r}", lpw, "0")
    put(sw, f"D{r}", f"=B{r}/C{r}", fmt="0.00"); put(sw, f"E{r}", f"=D{r}/10", fmt="0.0000"); put(sw, f"F{r}", f"=D{r}/1000", fmt="0.00000")
    inp(sw, f"G{r}", der, "0.000"); inp(sw, f"H{r}", ho, "0.0"); inp(sw, f"I{r}", hc, "0.0"); inp(sw, f"J{r}", hout, "0.0")
    put(sw, f"K{r}", f"=IF($B{r}<=1000,$E$22/1000*($E{r}/$E$5),$F$22*$E{r}*$B$22)", fmt="0.0000")
    put(sw, f"L{r}", f"=$D$23*($E{r}/100)*$G{r}", fmt="0.0000")
    put(sw, f"M{r}", f"=$D$24*($E{r}/100)*$G{r}", fmt="0.0000")
put(sw, "A12", "Reference: 1 full sun", bold=True); put(sw, "B12", 120000, fmt="#,##0"); put(sw, "C12", 120, fmt="0")
put(sw, "D12", "=B12/C12", fmt="0"); put(sw, "E12", "=D12/10", fmt="0.0"); put(sw, "F12", "=D12/1000", fmt="0.00")
put(sw, "A13", "Hours per day at each lux band are ESTIMATES from the research report; the rest of the day counts as zero light. Derating = mono-Si efficiency at that light divided by its 1-sun efficiency (UNVERIFIED snippets).", F_NOTE, wrap=True)
sw.merge_cells("A13:M13"); sw.row_dimensions[13].height = 28

head(sw, 15, ["Profile", "Fraction of the time the face is uncovered by a sleeve"])
for i, (n, f) in enumerate((("Office worker", 0.7), ("Commuter", 0.6), ("Outdoor enthusiast", 0.8))):
    put(sw, f"A{16 + i}", n); inp(sw, f"B{16 + i}", f, "0.00")
put(sw, "C16", "ESTIMATES: no published source.", F_NOTE)

head(sw, 21, ["Cell", "Footprint area (cm2)", "Type", "Power at 1 sun (mW)", "Measured at 200 lux (uW)", "Outdoor efficiency (a-Si only)", "Size (mm)", "Status"])
sw.row_dimensions[21].height = 40
cells = [("AM-1606C", "=1.5*1.5", "a-Si (glass)", None, 8.84, 0.045, "15 x 15 x 0.7", "200 lux VERIFIED (Panasonic catalog: 2.6 V x 3.4 uA). Outdoor efficiency 4.5% is an ESTIMATE (similar Amorton outdoor cells run 3.6-5.4%)."),
         ("KXOB25-07X1TF", "=1.135*0.8", "mono-Si", 14.3, None, None, "11.35 x 8 x 1.2", "1-sun power VERIFIED (Anysolar datasheet, preliminary). Low-light power is an ESTIMATE."),
         ("KXOB081K06TF (yours)", "=2.3*2.5", "mono-Si", 105.9, None, None, "23 x 25 x 1.2", "1-sun power VERIFIED (Anysolar datasheet, preliminary). Low-light power is an ESTIMATE.")]
for i, (n, a, t, p1, p200, eff, size, st) in enumerate(cells):
    r = 22 + i
    put(sw, f"A{r}", "Panasonic " + n if i == 0 else n, bold=True)
    put(sw, f"B{r}", a, fmt="0.00"); put(sw, f"C{r}", t)
    if p1 is not None: inp(sw, f"D{r}", p1, "0.0")
    if p200 is not None: inp(sw, f"E{r}", p200, "0.00")
    if eff is not None: inp(sw, f"F{r}", eff, "0.0%")
    put(sw, f"G{r}", size); put(sw, f"H{r}", st, F_NOTE)
# fix the a-Si row's name to match the dropdown (short name)
sw["A22"].value = "AM-1606C"

head(sw, 27, ["Harvest per day (mWh)", "Office worker", "Commuter", "Outdoor enthusiast"])
power_cols = {"AM-1606C": "K", "KXOB25-07X1TF": "L", "KXOB081K06TF (yours)": "M"}
hour_cols = ["H", "I", "J"]
for i, name in enumerate(("AM-1606C", "KXOB25-07X1TF", "KXOB081K06TF (yours)")):
    r = 28 + i
    put(sw, f"A{r}", name, bold=True)
    for j, hc in enumerate(hour_cols):
        col = "BCD"[j]
        put(sw, f"{col}{r}", f"=SUMPRODUCT(${hc}$5:${hc}$11,${power_cols[name]}$5:${power_cols[name]}$11)*Inputs!$B$39*$B${16 + j}*Inputs!$B$40*Inputs!$B$41*Inputs!$B$42", fmt="0.00", fill=FILL_OUT)
put(sw, "A32", "Harvest = sum over light bands of (hours x cell power) x wrist x sleeve x optical x harvester x charge efficiency. Compare with the daily load on the 'Power budget' tab.", F_NOTE, wrap=True)
sw.merge_cells("A32:M32"); sw.row_dimensions[32].height = 28
for col, w in zip("ABCDEFGHIJKLM", (30, 16, 14, 14, 16, 16, 18, 14, 14, 14, 16, 18, 18)):
    sw.column_dimensions[col].width = w
dv2 = DataValidation(type="list", formula1="=Solar!$A$28:$A$30", allow_blank=False); ws.add_data_validation(dv2); dv2.add("B37")
dv3 = DataValidation(type="list", formula1="=Solar!$B$27:$D$27", allow_blank=False); ws.add_data_validation(dv3); dv3.add("B38")

# =============================================================== Power budget
wp = wb.create_sheet("Power budget")
wp.column_dimensions["A"].width = 52
for col in "BCDE":
    wp.column_dimensions[col].width = 20
put(wp, "A1", "Power budget: what the watch uses and what solar could give back", F_TITLE)
section(wp, 3, "A. While the display is on (right now, with the LEDs chosen on the Inputs tab)", 5)
put(wp, "A4", "LED current drawn from the battery (mA)"); put(wp, "B4", "=SUM('Every LED'!L5:L136)", fmt="0.000")
put(wp, "A5", "LED power (mW)"); put(wp, "B5", "=SUM('Every LED'!M5:M136)", fmt="0.000")
put(wp, "A6", "MCU power while awake (mW)"); put(wp, "B6", "=Inputs!B5*Inputs!B27", fmt="0.000")
put(wp, "A7", "TOTAL display power (mW)", bold=True); put(wp, "B7", "=B5+B6", fmt="0.000", fill=FILL_OUT)
put(wp, "A8", "Energy for one wake (mWh)"); put(wp, "B8", "=B7*Inputs!B25/3600", fmt="0.00000")
put(wp, "A9", "Display energy per day (mWh)", bold=True); put(wp, "B9", "=B8*Inputs!B26", fmt="0.000", fill=FILL_OUT)

section(wp, 11, "B. Always on (waiting)", 5)
head(wp, 12, ["Component", "Current (uA)", "Supply (V)", "Power (uW)", "Energy per day (mWh)"])
put(wp, "A13", "MCU deep sleep + clock"); put(wp, "B13", "=Inputs!B30", F_LINK, "0.00"); put(wp, "C13", "=Inputs!B5", F_LINK, "0.00")
put(wp, "A14", "3.3 V regulator quiescent"); put(wp, "B14", "=Inputs!B31", F_LINK, "0.00"); put(wp, "C14", "=Inputs!B5", F_LINK, "0.00")
put(wp, "A15", "=\"Wake method: \"&Inputs!B32")
put(wp, "B15", "=INDEX(Lists!B2:B6,MATCH(Inputs!B32,Lists!A2:A6,0))", fmt="0.00"); put(wp, "C15", "=INDEX(Lists!C2:C6,MATCH(Inputs!B32,Lists!A2:A6,0))", fmt="0.00")
for r in (13, 14, 15):
    put(wp, f"D{r}", f"=B{r}*C{r}", fmt="0.0"); put(wp, f"E{r}", f"=D{r}*24/1000", fmt="0.000")
put(wp, "A16", "Battery self-discharge"); put(wp, "E16", "=Inputs!B7*Inputs!B8/30", fmt="0.000"); put(wp, "D16", "=E16*1000/24", fmt="0.0")
put(wp, "A17", "TOTAL always-on", bold=True); put(wp, "D17", "=SUM(D13:D16)", fmt="0.0", fill=FILL_OUT); put(wp, "E17", "=SUM(E13:E16)", fmt="0.000", fill=FILL_OUT)

section(wp, 19, "C. Daily total and battery life with no solar", 5)
put(wp, "A20", "Display energy per day (mWh)"); put(wp, "B20", "=B9", fmt="0.000")
put(wp, "A21", "Always-on energy per day (mWh)"); put(wp, "B21", "=E17", fmt="0.000")
put(wp, "A22", "TOTAL load per day (mWh)", bold=True); put(wp, "B22", "=B20+B21", fmt="0.000", fill=FILL_OUT)
put(wp, "A23", "Battery life with no solar (days)", bold=True); put(wp, "B23", "=IF(B22>0,Inputs!B7/B22,0)", fmt="0", fill=FILL_OUT)
put(wp, "A24", "Share of the load that is the display"); put(wp, "B24", "=IF(B22>0,B20/B22,0)", fmt="0%")
put(wp, "A25", "Share of the load that is the wake sensor"); put(wp, "B25", "=IF(B22>0,E15/B22,0)", fmt="0%")

section(wp, 27, "D. With the solar cell and profile picked on the Inputs tab", 5)
put(wp, "A28", "Cell"); put(wp, "B28", "=Inputs!B37", F_LINK)
put(wp, "A29", "Profile"); put(wp, "B29", "=Inputs!B38", F_LINK)
put(wp, "A30", "Harvest per day (mWh)", bold=True); put(wp, "B30", "=INDEX(Solar!B28:D30,MATCH(Inputs!B37,Solar!A28:A30,0),MATCH(Inputs!B38,Solar!B27:D27,0))", fmt="0.000", fill=FILL_OUT)
put(wp, "A31", "Solar covers this much of the load"); put(wp, "B31", "=IF(B22>0,B30/B22,0)", fmt="0%")
put(wp, "A32", "Net drain on the battery per day (mWh)"); put(wp, "B32", "=MAX(0,B22-B30)", fmt="0.000")
put(wp, "A33", "Battery life with solar (days)", bold=True); put(wp, "B33", '=IF(B32<=0.0001,"never runs down",Inputs!B7/B32)', fmt="0", fill=FILL_OUT)

section(wp, 35, "E. What if I change the PWM dimming? (all three LEDs at the same duty; other inputs as above)", 5)
head(wp, 36, ["PWM duty (all LEDs)", "Total current of the 3 lit LEDs (mA)", "Display energy per day (mWh)", "Total load per day (mWh)", "Battery life with no solar (days)"])
wp.row_dimensions[36].height = 42
for i, d in enumerate((0.05, 0.1, 0.2, 0.3, 0.5, 0.75, 1.0)):
    r = 37 + i
    inp(wp, f"A{r}", d, "0%")
    put(wp, f"B{r}", f"=('LED types'!$F$5+'LED types'!$F$6+'LED types'!$F$7)*'LED types'!$G$5*A{r}", fmt="0.000")
    put(wp, f"C{r}", f"=(Inputs!$B$5*B{r}+Inputs!$B$5*Inputs!$B$27)*Inputs!$B$25/3600*Inputs!$B$26", fmt="0.000")
    put(wp, f"D{r}", f"=C{r}+$E$17", fmt="0.000")
    put(wp, f"E{r}", f"=IF(D{r}>0,Inputs!$B$7/D{r},0)", fmt="0")

section(wp, 46, "F. Where the energy goes (mWh per day)", 5)
labels = [("Display", "=B9"), ("MCU sleep", "=E13"), ("Regulator", "=E14"), ("Wake sensor", "=E15"), ("Self-discharge", "=E16")]
for i, (n, f) in enumerate(labels):
    put(wp, f"A{47 + i}", n); put(wp, f"B{47 + i}", f, fmt="0.000")


section(wp, 54, "G. Special mode (fluid / gravity / tilt game)", 5)
put(wp, "A55", "LED on-current used (mA)"); put(wp, "B55", "=AVERAGE('LED types'!F5:F7)", fmt="0.000")
put(wp, "A56", "Average current per lit LED in special mode (mA)"); put(wp, "B56", "=B55/Inputs!B47*Inputs!B48", fmt="0.000")
put(wp, "A57", "Total LED current while playing (mA)"); put(wp, "B57", "=Inputs!B46*B56", fmt="0.00")
put(wp, "A58", "Motion sensor current in game mode (uA)"); put(wp, "B58", "=INDEX(Lists!G2:G6,MATCH(Inputs!B51,Lists!F2:F6,0))", fmt="0.0")
put(wp, "A59", "Total power while playing (mW)", bold=True); put(wp, "B59", "=Inputs!B5*(B57+Inputs!B53)+B58*3.3/1000", fmt="0.00", fill=FILL_OUT)
put(wp, "A60", "Special-mode energy per day (mWh)", bold=True); put(wp, "B60", "=B59*Inputs!B45/60", fmt="0.000", fill=FILL_OUT)
put(wp, "A61", "TOTAL load per day including special mode (mWh)", bold=True); put(wp, "B61", "=B22+B60", fmt="0.000", fill=FILL_OUT)
put(wp, "A62", "Battery life including special mode, no solar (days)", bold=True); put(wp, "B62", "=IF(B61>0,Inputs!B7/B61,0)", fmt="0", fill=FILL_OUT)
put(wp, "A63", "Share of the day's energy that is special mode"); put(wp, "B63", "=IF(B61>0,B60/B61,0)", fmt="0%")
put(wp, "A64", "Worst-case current through one MCU ROW pin (mA)"); put(wp, "B64", "=Inputs!B49*B55", fmt="0.0")
put(wp, "A65", "Limit you set for one pin (mA)"); put(wp, "B65", "=Inputs!B50", F_LINK, "0.0")
put(wp, "A66", "Verdict on driving the rows straight from the MCU", bold=True)
put(wp, "B66", '=IF(B64<=B65,"OK as is","TOO HIGH: row pins need driver transistors, or a lower on-current / fewer LEDs per row")', fill=FILL_OUT)
wp.merge_cells("B66:E66")

ch = LineChart(); ch.title = "Battery life vs PWM duty (no solar)"; ch.y_axis.title = "days"; ch.x_axis.title = "PWM duty"
ch.add_data(Reference(wp, min_col=5, min_row=36, max_row=43), titles_from_data=True); ch.set_categories(Reference(wp, min_col=1, min_row=37, max_row=43))
ch.height = 8; ch.width = 15; wp.add_chart(ch, "G3")
bc = BarChart(); bc.type = "col"; bc.title = "Daily energy by consumer (mWh)"
bc.add_data(Reference(wp, min_col=2, min_row=47, max_row=51), titles_from_data=False); bc.set_categories(Reference(wp, min_col=1, min_row=47, max_row=51))
bc.legend = None; bc.height = 8; bc.width = 15; wp.add_chart(bc, "G20")

# =============================================================== Notes
wn = wb.create_sheet("Notes")
wn.column_dimensions["A"].width = 130
notes = [
    ("What this workbook is", True),
    ("A live model of the single-colour LED watch. Change the yellow cells on 'Inputs' (or the LED/solar tables) and every other tab recalculates.", False),
    ("", False),
    ("Tabs", True),
    ("Inputs: all the levers. LED types: electrical model per LED size. Every LED: all 132 LEDs, where each sits in the matrix, and its power draw. Solar: light units, cell power, daily harvest. Power budget: results and what-if tables.", False),
    ("", False),
    ("What 'Variant A' and 'Variant B' mean (from the schematics)", True),
    ("Variant A = the battery is charged through two small pads (like a charging dock or a USB cable): a simple charger chip (MCP73831) fills the LiPo when you plug it in. No solar.", False),
    ("Variant B = the battery is charged by a solar cell on the face, through an energy-harvesting chip (BQ25570). You could still add pads for backup.", False),
    ("This workbook works for either: the 'always-on' and 'display' numbers are the same; only the Solar section matters for Variant B.", False),
    ("", False),
    ("How PWM and multiplexing relate", True),
    ("Multiplexing = switching rows on one at a time so a few pins can drive many LEDs. PWM = switching an LED on for only part of its slot to make it dimmer. Both are just timers in firmware. Brightness and energy both follow the average current.", False),
    ("", False),
    ("Special mode (section 7 of Inputs, section G of Power budget)", True),
    ("A fluid/gravity or tilt game only needs the direction of gravity, so an ACCELEROMETER is enough (a gyroscope is not needed). The model charges you for the LEDs, the MCU and the sensor for the minutes you choose, and checks whether one MCU row pin can carry all the LEDs lit in that row.", False),
    ("", False),
    ("Honesty notes", True),
    ("Tags: VERIFIED = taken from a manufacturer datasheet or catalog; ESTIMATE = an assumption of mine. Estimates are in the input cells so you can overwrite them. The LED current at low drive depends strongly on Vf, which is a place the real number matters: measure it on a prototype.", False),
    ("Solar cell low-light power is not published for Anysolar cells; those figures are derived from one set of unverified efficiency snippets. Measure before relying on them.", False),
    ("Sources: see solar-watch-evidence.md, solar-power-research.md, solar-cell-sourcing.md, mcu-research.md in the project folder.", False),
]
for i, (t, b) in enumerate(notes, start=1):
    c = wn.cell(row=i, column=1, value=t); c.font = F_BOLD if b else F_NORM; c.alignment = Alignment(wrap_text=True, vertical="top")

exec(open("build_spec_tabs.py", encoding="utf-8").read())
exec(open("build_layout_tabs.py", encoding="utf-8").read())

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save("Watch-Power-Calculator.xlsx")
print("saved")
